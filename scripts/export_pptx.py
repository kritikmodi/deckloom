#!/usr/bin/env python3
"""Export the built deck to .pptx, one rendered slide per page.

Be clear about what this is: each slide becomes a full-bleed image inside a
PowerPoint file. It opens anywhere, presents correctly, and survives being
emailed to someone who only has PowerPoint. The text is not editable in
PowerPoint, because the layout was done by a browser.

That is usually the right trade. If a colleague needs to edit wording, they
should edit content.json and rebuild, which is the entire point of the pipeline.
If they genuinely need native editable shapes, this is the wrong tool and you
want a real .pptx authoring library instead.

    python3 scripts/export_pptx.py                    # out/deck.pdf -> out/deck.pptx
    python3 scripts/export_pptx.py --dpi 220          # sharper, larger file
"""
import argparse, pathlib, sys

HERE = pathlib.Path(__file__).resolve().parent.parent


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--pdf", default="out/deck.pdf")
    ap.add_argument("--out", default="")
    ap.add_argument("--dpi", type=int, default=180,
                    help="180 is crisp on a projector; 300 for print handouts")
    a = ap.parse_args()

    try:
        import pymupdf
    except ImportError:
        sys.exit("needs pymupdf:  pip install pymupdf")
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        sys.exit("needs python-pptx:  pip install python-pptx")

    pdf = HERE / a.pdf
    if not pdf.exists():
        sys.exit(f"no deck at {pdf}. Run scripts/build.py first.")
    out = HERE / (a.out or str(pdf.with_suffix(".pptx")))

    doc = pymupdf.open(pdf)
    page = doc[0]
    prs = Presentation()
    # Match the PDF page exactly, so nothing is letterboxed or cropped.
    prs.slide_width = Emu(int(page.rect.width / 72 * 914400))
    prs.slide_height = Emu(int(page.rect.height / 72 * 914400))
    blank = prs.slide_layouts[6]          # 6 is the blank layout

    n = len(doc)
    tmp = out.parent / "_pptx_tmp"
    tmp.mkdir(parents=True, exist_ok=True)
    try:
        for i, pg in enumerate(doc, start=1):
            img = tmp / f"p{i:03d}.png"
            pg.get_pixmap(dpi=a.dpi).save(img)
            s = prs.slides.add_slide(blank)
            s.shapes.add_picture(str(img), 0, 0,
                                 width=prs.slide_width, height=prs.slide_height)
        prs.save(out)
    finally:
        for f in tmp.glob("*.png"):
            f.unlink()
        tmp.rmdir()
    doc.close()

    mb = out.stat().st_size / 1e6
    shown = out.relative_to(HERE) if out.is_relative_to(HERE) else out
    print(f"wrote {shown}  ({n} slides, {mb:.1f} MB, {a.dpi} dpi)")
    print("slides are images: to change wording, edit content.json and rebuild")


if __name__ == "__main__":
    main()
